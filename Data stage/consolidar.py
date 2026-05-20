"""
Data Stage SAT → Excel Multi-Mes
Consolida 2 o más zips de Data Stage SAT en un solo archivo Excel.
Cada hoja contiene los datos de todos los meses apilados (zip1 → zip2 → zip3...).
Columna A = Pedimento_Unificado (AA-AAA-PPPP-PPPPPPP) en todas las hojas.

Formato Pedimento_Unificado:
  AA       = año (2 dígitos, leído de Fecha_Inicial en Resumen.asc de cada zip)
  AAA      = SeccionAduanera (3 dígitos)
  PPPP     = Patente (4 dígitos)
  PPPPPPP  = Pedimento (7 dígitos)

Uso:
    python consolidar.py <salida.xlsx> <zip1> <zip2> [<zip3> ...]
"""

import pandas as pd
from openpyxl import Workbook
import os, sys, re, zipfile, shutil, tempfile
import numpy as np
from collections import defaultdict


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def find_asc_dir(base_dir):
    asc_files = []
    for root, dirs, files in os.walk(base_dir):
        dirs[:] = [d for d in dirs if d != '__MACOSX']
        for f in files:
            if f.endswith('.asc'):
                asc_files.append(os.path.join(root, f))
    if not asc_files:
        raise FileNotFoundError(f"No se encontraron .asc en {base_dir}")
    dirs_found = set(os.path.dirname(f) for f in asc_files)
    if len(dirs_found) > 1:
        from collections import Counter
        return Counter(os.path.dirname(f) for f in asc_files).most_common(1)[0][0]
    return list(dirs_found)[0]


def get_folio(asc_dir):
    files = [f for f in os.listdir(asc_dir) if f.endswith('.asc')]
    if not files:
        return "DATASTAGE"
    parts = files[0].split('_')
    return parts[0] if len(parts) >= 2 else "DATASTAGE"


def get_anio_from_resumen(asc_dir, folio):
    """
    Lee Resumen.asc y extrae el año (2 dígitos) de Fecha_Inicial.
    Ejemplo: '2026-03-01' → '26'
    Fallback: año actual.
    """
    resumen_path = os.path.join(asc_dir, f"{folio}_Resumen.asc")
    if not os.path.isfile(resumen_path):
        for f in os.listdir(asc_dir):
            if 'resumen' in f.lower() and f.endswith('.asc'):
                resumen_path = os.path.join(asc_dir, f)
                break
        else:
            import datetime
            return str(datetime.datetime.now().year)[-2:]
    try:
        df = pd.read_csv(resumen_path, sep='|', dtype=str, encoding='latin-1', index_col=False)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        col_map = {c.strip().lower(): c for c in df.columns}
        fecha_col = col_map.get('fecha_inicial')
        if fecha_col and len(df) > 0:
            fecha_val = str(df[fecha_col].iloc[0]).strip()
            match = re.match(r'(\d{4})', fecha_val)
            if match:
                return match.group(1)[-2:]
    except Exception:
        pass
    import datetime
    return str(datetime.datetime.now().year)[-2:]


def read_asc(path):
    df = pd.read_csv(
        path,
        sep='|',
        dtype=str,
        encoding='latin-1',
        index_col=False  # CRÍTICO: evita bug de desfase
    )
    df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
    return df


def build_pedimento_unificado(df, anio_zip):
    """
    AA-AAA-PPPP-PPPPPPP
    AA      = anio_zip (del Resumen del zip)
    AAA     = SeccionAduanera
    PPPP    = Patente
    PPPPPPP = Pedimento
    """
    col_map = {c.strip().lower(): c for c in df.columns}

    def pad(col_name, width):
        real_col = col_map.get(col_name.lower()) if col_name else None
        if real_col is None or real_col not in df.columns:
            return pd.Series(['0' * width] * len(df), index=df.index)
        s = df[real_col].fillna('').astype(str).str.strip()
        s = s.str.replace(r'\.0$', '', regex=True)
        return s.str.zfill(width)

    anio      = pd.Series([anio_zip] * len(df), index=df.index)
    aduana    = pad('seccionaduanera', 3)
    patente   = pad('patente',         4)
    pedimento = pad('pedimento',       7)

    return anio + '-' + aduana + '-' + patente + '-' + pedimento


# ---------------------------------------------------------------------------
# Leer todos los zips y agrupar DataFrames por hoja
# ---------------------------------------------------------------------------

def extract_zip(zip_path, tmp_root):
    folder_name = os.path.splitext(os.path.basename(zip_path))[0]
    dest = os.path.join(tmp_root, folder_name)
    os.makedirs(dest, exist_ok=True)
    with zipfile.ZipFile(zip_path, 'r') as z:
        z.extractall(dest)
    if not any(f.endswith('.asc') for f in os.listdir(dest)):
        dest = find_asc_dir(dest)
    return dest


def load_all_zips(zip_paths, tmp_root):
    """
    Retorna: { sheet_name: [(df, anio_zip), ...] }  en orden de zip_paths
    """
    sheets = defaultdict(list)

    for zip_path in zip_paths:
        print(f"\n📦 Procesando: {os.path.basename(zip_path)}")
        asc_dir  = extract_zip(zip_path, tmp_root)
        folio    = get_folio(asc_dir)
        anio_zip = get_anio_from_resumen(asc_dir, folio)
        files    = sorted([f for f in os.listdir(asc_dir) if f.endswith('.asc')])
        print(f"   Folio: {folio} | Año: {anio_zip} | Archivos .asc: {len(files)}")

        for fname in files:
            sheet_name = fname.replace(f"{folio}_", "").replace(".asc", "")
            df = read_asc(os.path.join(asc_dir, fname))
            sheets[sheet_name].append((df, anio_zip))
            print(f"   {sheet_name}: {len(df)} filas")

    return sheets


# ---------------------------------------------------------------------------
# Consolidar y escribir Excel
# ---------------------------------------------------------------------------

def write_consolidated_sheet(ws, df_combined, anio_combined, fallback_columns=None):
    """
    df_combined  : DataFrame con todos los meses apilados
    anio_combined: Series con el año correcto por fila (uno por cada mes)
    """
    # Siempre escribir header, aunque no haya filas de datos
    columns_to_use = list(df_combined.columns) if not df_combined.empty else (fallback_columns or [])
    all_columns = ['Pedimento_Unificado'] + columns_to_use

    for ci, col in enumerate(all_columns, 1):
        ws.cell(row=1, column=ci, value=col)

    # Filtros en el header aunque no haya datos
    if len(all_columns) > 0:
        last_col = ws.cell(row=1, column=len(all_columns)).column_letter
        ws.auto_filter.ref = f"A1:{last_col}1"

    if df_combined.empty:
        return

    pu_series   = build_pedimento_unificado_combined(df_combined, anio_combined)
    total_cols = len(all_columns)

    # Datos
    pu_values = pu_series.tolist()
    df_values = df_combined.values

    for ri in range(len(df_values)):
        ws.cell(row=ri + 2, column=1, value=pu_values[ri])
        for ci, val in enumerate(df_values[ri], 2):
            if val is None or (isinstance(val, float) and np.isnan(val)) or str(val) == 'nan':
                ws.cell(row=ri + 2, column=ci, value=None)
            else:
                ws.cell(row=ri + 2, column=ci, value=val)

    # (header y filtros ya escritos arriba)


def build_pedimento_unificado_combined(df, anio_series):
    """
    Versión para multi-mes: anio_series es una pd.Series con el año
    correcto para cada fila (ya que distintas filas pueden ser de distintos meses).
    """
    col_map = {c.strip().lower(): c for c in df.columns}

    def pad(col_name, width):
        real_col = col_map.get(col_name.lower()) if col_name else None
        if real_col is None or real_col not in df.columns:
            return pd.Series(['0' * width] * len(df), index=df.index)
        s = df[real_col].fillna('').astype(str).str.strip()
        s = s.str.replace(r'\.0$', '', regex=True)
        return s.str.zfill(width)

    aduana    = pad('seccionaduanera', 3)
    patente   = pad('patente',         4)
    pedimento = pad('pedimento',       7)

    return anio_series.reset_index(drop=True) + '-' + \
           aduana.reset_index(drop=True)      + '-' + \
           patente.reset_index(drop=True)     + '-' + \
           pedimento.reset_index(drop=True)


def consolidate(zip_paths, output_path):
    tmp_root = tempfile.mkdtemp(prefix='datastage_multi_')
    try:
        sheets = load_all_zips(zip_paths, tmp_root)

        print(f"\n🔗 Consolidando {len(sheets)} tipos de hoja...")
        wb = Workbook()
        wb.remove(wb.active)

        sheet_order = sorted(sheets.keys(), key=lambda x: (
            int(x) if x.isdigit() else float('inf'), x
        ))

        for sheet_name in sheet_order:
            entries = sheets[sheet_name]  # lista de (df, anio_zip)

            # Apilar DataFrames y construir Series de año correspondiente
            dfs   = []
            anios = []
            for df, anio_zip in entries:
                dfs.append(df)
                anios.append(pd.Series([anio_zip] * len(df)))

            df_combined    = pd.concat(dfs,   ignore_index=True) if dfs else pd.DataFrame()
            anio_combined  = pd.concat(anios, ignore_index=True) if anios else pd.Series(dtype=str)

            # Fallback columns: usar las del primer df de esta hoja (para hojas vacías)
            fallback_cols = list(entries[0][0].columns) if entries else []
            ws = wb.create_sheet(title=sheet_name)
            write_consolidated_sheet(ws, df_combined, anio_combined, fallback_columns=fallback_cols)
            print(f"  {sheet_name}: {len(df_combined)} filas totales ({len(entries)} mes(es))")

        wb.save(output_path)
        print(f"\n💾 Guardado: {output_path} | Hojas: {len(wb.sheetnames)} | Zips: {len(zip_paths)}")
        verify_output(output_path)

    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)


# ---------------------------------------------------------------------------
# Verificación
# ---------------------------------------------------------------------------

def verify_output(output_path):
    from openpyxl import load_workbook
    wb      = load_workbook(output_path)
    issues  = []
    pattern = re.compile(r'^\d{2}-\d{3}-\d{4}-\d{7}$')

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        if ws.max_row < 2:
            continue
        header_a = ws['A1'].value
        val_a2   = str(ws['A2'].value or '').strip()
        if header_a != 'Pedimento_Unificado':
            issues.append(f"  ERROR en '{sheet}': A1='{header_a}'")
        elif not pattern.match(val_a2):
            issues.append(f"  FORMATO en '{sheet}': A2='{val_a2}' (esperado: AA-AAA-PPPP-PPPPPPP)")

    if issues:
        print("⚠️  Advertencias:")
        for i in issues: print(i)
    else:
        print("✅ Verificación OK: Pedimento_Unificado (AA-AAA-PPPP-PPPPPPP) en todas las hojas con datos")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Uso: python consolidar.py <salida.xlsx> <zip1> <zip2> [<zip3> ...]")
        sys.exit(1)

    output_path = sys.argv[1]
    zip_paths   = sys.argv[2:]

    for z in zip_paths:
        if not os.path.isfile(z):
            print(f"❌ No encontrado: {z}")
            sys.exit(1)

    print(f"🗂  Zips a consolidar: {len(zip_paths)}")
    for i, z in enumerate(zip_paths, 1):
        print(f"  {i}. {os.path.basename(z)}")

    consolidate(zip_paths, output_path)
