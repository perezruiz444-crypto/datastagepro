"""
Data Stage SAT → Excel v2
Conversión de archivos .asc a Excel con filtros, sin desfase de columnas,
y con columna A = Pedimento_Unificado (AA-AAA-PPPP-PPPPPPP) en todas las hojas.

Formato: AA-AAA-PPPP-PPPPPPP
  AA       = año (últimos 2 dígitos de Fecha_Inicial en Resumen.asc)
  AAA      = SeccionAduanera (3 dígitos)
  PPPP     = Patente (4 dígitos)
  PPPPPPP  = Pedimento (7 dígitos)

Uso:
    python convert_script.py <directorio_con_asc> <archivo_salida.xlsx>
"""

import pandas as pd
from openpyxl import Workbook
import os, sys, re
import numpy as np


def find_asc_dir(base_dir):
    asc_files = []
    for root, dirs, files in os.walk(base_dir):
        dirs[:] = [d for d in dirs if d != '__MACOSX']
        for f in files:
            if f.endswith('.asc'):
                asc_files.append(os.path.join(root, f))
    if not asc_files:
        raise FileNotFoundError(f"No se encontraron .asc en {base_dir}")
    dirs_found = set(os.path.dirname(f) for f in asc_files)
    if len(dirs_found) > 1:
        from collections import Counter
        return Counter(os.path.dirname(f) for f in asc_files).most_common(1)[0][0]
    return list(dirs_found)[0]


def get_folio(asc_dir):
    files = [f for f in os.listdir(asc_dir) if f.endswith('.asc')]
    if not files:
        return "DATASTAGE"
    parts = files[0].split('_')
    return parts[0] if len(parts) >= 2 else "DATASTAGE"


def get_anio_from_resumen(asc_dir, folio):
    """
    Lee Resumen.asc y extrae el año (2 dígitos) de Fecha_Inicial.
    Ejemplo: '2026-03-01' → '26'
    Fallback: año actual si no se puede leer.
    """
    resumen_path = os.path.join(asc_dir, f"{folio}_Resumen.asc")
    if not os.path.isfile(resumen_path):
        # Buscar case-insensitive
        for f in os.listdir(asc_dir):
            if 'resumen' in f.lower() and f.endswith('.asc'):
                resumen_path = os.path.join(asc_dir, f)
                break
        else:
            import datetime
            return str(datetime.datetime.now().year)[-2:]

    try:
        df = pd.read_csv(resumen_path, sep='|', dtype=str, encoding='latin-1', index_col=False)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        # Buscar columna Fecha_Inicial (case-insensitive)
        col_map = {c.strip().lower(): c for c in df.columns}
        fecha_col = col_map.get('fecha_inicial')
        if fecha_col and len(df) > 0:
            fecha_val = str(df[fecha_col].iloc[0]).strip()
            # Formato esperado: YYYY-MM-DD o YYYY/MM/DD
            match = re.match(r'(\d{4})', fecha_val)
            if match:
                return match.group(1)[-2:]  # últimos 2 dígitos del año
    except Exception:
        pass

    import datetime
    return str(datetime.datetime.now().year)[-2:]


def read_asc(path):
    df = pd.read_csv(
        path,
        sep='|',
        dtype=str,
        encoding='latin-1',
        index_col=False  # CRÍTICO: evita bug de desfase
    )
    df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
    return df


def build_pedimento_unificado(df, anio_zip):
    """
    Construye Pedimento_Unificado con formato AA-AAA-PPPP-PPPPPPP.

    AA      = anio_zip  (año leído del Resumen del zip, 2 dígitos)
    AAA     = SeccionAduanera (3 dígitos)
    PPPP    = Patente (4 dígitos)
    PPPPPPP = Pedimento (7 dígitos)

    Si un campo no existe en la hoja, ese segmento queda como ceros.
    """
    col_map = {c.strip().lower(): c for c in df.columns}

    def pad(col_name, width):
        real_col = col_map.get(col_name.lower()) if col_name else None
        if real_col is None or real_col not in df.columns:
            return pd.Series(['0' * width] * len(df), index=df.index)
        s = df[real_col].fillna('').astype(str).str.strip()
        s = s.str.replace(r'\.0$', '', regex=True)
        return s.str.zfill(width)

    anio      = pd.Series([anio_zip] * len(df), index=df.index)
    aduana    = pad('seccionaduanera', 3)
    patente   = pad('patente',         4)
    pedimento = pad('pedimento',       7)

    return anio + '-' + aduana + '-' + patente + '-' + pedimento


def write_sheet(ws, df, anio_zip):
    pu_series   = build_pedimento_unificado(df, anio_zip)
    all_columns = ['Pedimento_Unificado'] + list(df.columns)
    total_cols  = len(all_columns)

    # Header
    for ci, col in enumerate(all_columns, 1):
        ws.cell(row=1, column=ci, value=col)

    # Datos
    pu_values = pu_series.tolist()
    df_values = df.values

    for ri in range(len(df_values)):
        ws.cell(row=ri + 2, column=1, value=pu_values[ri])
        for ci, val in enumerate(df_values[ri], 2):
            if val is None or (isinstance(val, float) and np.isnan(val)) or str(val) == 'nan':
                ws.cell(row=ri + 2, column=ci, value=None)
            else:
                ws.cell(row=ri + 2, column=ci, value=val)

    # Filtros automáticos
    if total_cols > 0:
        last_col = ws.cell(row=1, column=total_cols).column_letter
        ws.auto_filter.ref = f"A1:{last_col}1"


def verify_output(output_path):
    from openpyxl import load_workbook
    wb      = load_workbook(output_path)
    issues  = []
    pattern = re.compile(r'^\d{2}-\d{3}-\d{4}-\d{7}$')

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        if ws.max_row < 2:
            continue
        header_a = ws['A1'].value
        val_a2   = str(ws['A2'].value or '').strip()
        if header_a != 'Pedimento_Unificado':
            issues.append(f"  ERROR en '{sheet}': A1='{header_a}' (esperado: Pedimento_Unificado)")
        elif not pattern.match(val_a2):
            issues.append(f"  FORMATO en '{sheet}': A2='{val_a2}' (esperado: AA-AAA-PPPP-PPPPPPP)")

    if issues:
        print("⚠️  Advertencias:")
        for i in issues: print(i)
    else:
        print("✅ Verificación OK: Pedimento_Unificado (AA-AAA-PPPP-PPPPPPP) en todas las hojas con datos")


def convert(asc_dir, output_path):
    files = sorted([f for f in os.listdir(asc_dir) if f.endswith('.asc')])
    if not files:
        raise FileNotFoundError(f"No hay .asc en {asc_dir}")

    folio    = get_folio(asc_dir)
    anio_zip = get_anio_from_resumen(asc_dir, folio)
    print(f"Folio: {folio} | Año: {anio_zip} | Archivos .asc: {len(files)}")

    wb = Workbook()
    wb.remove(wb.active)

    for fname in files:
        path       = os.path.join(asc_dir, fname)
        sheet_name = fname.replace(f"{folio}_", "").replace(".asc", "")
        df         = read_asc(path)
        ws         = wb.create_sheet(title=sheet_name)
        write_sheet(ws, df, anio_zip)
        print(f"  {sheet_name}: {len(df)} filas, {len(df.columns)+1} cols (A=Pedimento_Unificado)")

    wb.save(output_path)
    print(f"\n💾 Guardado: {output_path} | Hojas: {len(wb.sheetnames)}")
    verify_output(output_path)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python convert_script.py <directorio_asc> <salida.xlsx>")
        sys.exit(1)

    asc_dir     = sys.argv[1]
    output_path = sys.argv[2]

    if not any(f.endswith('.asc') for f in os.listdir(asc_dir)):
        print(f"Buscando .asc dentro de {asc_dir}...")
        asc_dir = find_asc_dir(asc_dir)
        print(f"Directorio .asc encontrado: {asc_dir}")

    convert(asc_dir, output_path)
